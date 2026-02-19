import io
import os

try:
    from pptx import Presentation
    from pptx.util import Pt, Inches, Emu
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
    from pptx.util import Inches, Pt
except ImportError:
    raise ImportError("Library 'python-pptx' not found. Please run `pip install python-pptx`.")

from src.ai_engine import generate_narrative_content

# ─── BRAND PALETTE ─────────────────────────────────────────────────────────────
C_DARK    = RGBColor(0x0D, 0x1B, 0x2A)   # Deep navy  – slide bg / title bar
C_ACCENT  = RGBColor(0x00, 0xB0, 0xF0)   # Bright cyan – headings / accents
C_MID     = RGBColor(0x1A, 0x3C, 0x5E)   # Mid blue   – subtitle bars
C_LIGHT   = RGBColor(0xF0, 0xF4, 0xF8)   # Off-white  – body bg / text fields
C_WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
C_ORANGE  = RGBColor(0xFF, 0x6B, 0x35)   # Highlight / KPI callouts
C_GANTT   = RGBColor(0x00, 0xB0, 0xF0)   # Timeline fill cells

SLIDE_W = Inches(13.33)
SLIDE_H = Inches(7.5)

# ─── LOW-LEVEL HELPERS ─────────────────────────────────────────────────────────

def _rgb(r, g, b): return RGBColor(r, g, b)

def _add_rect(slide, left, top, width, height, fill_color, line_color=None):
    shape = slide.shapes.add_shape(1, left, top, width, height)  # MSO_SHAPE.RECTANGLE
    shape.fill.solid()
    shape.fill.fore_color.rgb = fill_color
    if line_color:
        shape.line.color.rgb = line_color
    else:
        shape.line.fill.background()
    return shape

def _add_text(slide, text, left, top, width, height,
              font_size=18, bold=False, color=C_WHITE,
              align=PP_ALIGN.LEFT, wrap=True, font_name="Calibri"):
    txb = slide.shapes.add_textbox(left, top, width, height)
    tf = txb.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.color.rgb = color
    run.font.name = font_name
    return txb

def _add_image_safe(slide, img_source, left, top, width=None, height=None):
    """Add image from file path or file-like / UploadedFile object."""
    if not img_source:
        return
    try:
        if isinstance(img_source, str):
            if not os.path.exists(img_source):
                return
            img_data = img_source
        elif hasattr(img_source, 'getbuffer'):
            img_data = io.BytesIO(img_source.getbuffer())
        else:
            img_data = img_source
        slide.shapes.add_picture(img_data, left, top, width=width, height=height)
    except Exception as e:
        print(f"Image insert error: {e}")

# ─── SLIDE BACKGROUNDS ─────────────────────────────────────────────────────────

def _set_dark_bg(slide):
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, C_DARK)

def _set_light_bg(slide):
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, C_LIGHT)
    # accent top bar
    _add_rect(slide, 0, 0, SLIDE_W, Inches(0.08), C_ACCENT)

def _add_section_heading(slide, title):
    """Adds a styled heading bar at top of content slides."""
    _add_rect(slide, 0, 0, SLIDE_W, Inches(1.1), C_MID)
    _add_text(slide, title,
              Inches(0.3), Inches(0.1), Inches(12.5), Inches(0.9),
              font_size=28, bold=True, color=C_ACCENT, align=PP_ALIGN.LEFT)

def _add_kpi_box(slide, left, top, width, height, label, value, bg=C_MID):
    _add_rect(slide, left, top, width, height, bg)
    _add_text(slide, label,
              left + Inches(0.1), top + Inches(0.05), width - Inches(0.2), Inches(0.35),
              font_size=11, bold=False, color=C_ACCENT)
    _add_text(slide, value,
              left + Inches(0.1), top + Inches(0.35), width - Inches(0.2), height - Inches(0.5),
              font_size=18, bold=True, color=C_WHITE)


# ─── SLIDE BUILDERS ────────────────────────────────────────────────────────────

def _slide_title(prs, data, proposal_type):
    """Slide 1: Title"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # Blank
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, C_DARK)
    # accent stripe
    _add_rect(slide, 0, Inches(3.0), SLIDE_W, Inches(0.06), C_ACCENT)

    client  = data.get("client_name", "Client")
    vendor  = data.get("vendor_name", "Oneture")
    timeline = data.get("delivery_timeline", "").replace("_", " ").title()
    title_map = {
        "Demand Forecasting": "AI-Powered Demand Forecasting",
        "Visual Inspection":  "AI-Powered Visual Inspection",
        "Chatbot":            "Enterprise AI Chatbot Solution",
    }
    title_text = title_map.get(proposal_type, proposal_type + " Proposal")
    _add_text(slide, f"PROPOSAL FOR {client.upper()}",
              Inches(0.6), Inches(1.2), Inches(10), Inches(0.6),
              font_size=14, bold=False, color=C_ACCENT)
    _add_text(slide, title_text,
              Inches(0.6), Inches(1.8), Inches(10), Inches(1.2),
              font_size=38, bold=True, color=C_WHITE)
    _add_text(slide, f"Prepared by {vendor}   |   {timeline}",
              Inches(0.6), Inches(3.2), Inches(10), Inches(0.5),
              font_size=16, bold=False, color=C_LIGHT)

    # logos
    _add_image_safe(slide, data.get("client_logo_path"),
                    Inches(0.5), Inches(5.8), height=Inches(1.2))
    _add_image_safe(slide, data.get("company_logo_path"),
                    Inches(11.3), Inches(5.8), height=Inches(1.2))


def _slide_agenda(prs, proposal_type):
    """Slide 2: Agenda"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Agenda")

    agenda_items = [
        ("01", "Executive Summary",          "High-level problem and proposed solution"),
        ("02", "Problem Statement",           "Current limitations and operational gaps"),
        ("03", "Data & Infrastructure",       "Data reality and technology posture"),
        ("04", "Proposed Solution",           "Intelligence layer and architecture"),
        ("05", "Security & Compliance",       "Standards, controls, and governance"),
        ("06", "Success Metrics",             "KPIs, validation, and acceptance criteria"),
        ("07", "Domain Deep Dive",            f"{proposal_type}-specific technical details"),
        ("08", "Implementation Timeline",    "Phased delivery plan"),
        ("09", "Commercial Summary",          "Budget and key assumptions"),
    ]
    col_w = Inches(6.2)
    for i, (num, title, desc) in enumerate(agenda_items):
        col = i % 2
        row = i // 2
        left = Inches(0.3) + col * col_w
        top  = Inches(1.3) + row * Inches(1.4)
        _add_rect(slide, left, top, col_w - Inches(0.2), Inches(1.2), C_MID)
        _add_text(slide, num,
                  left + Inches(0.1), top + Inches(0.05), Inches(0.5), Inches(0.5),
                  font_size=20, bold=True, color=C_ACCENT)
        _add_text(slide, title,
                  left + Inches(0.65), top + Inches(0.05), col_w - Inches(1.0), Inches(0.5),
                  font_size=15, bold=True, color=C_WHITE)
        _add_text(slide, desc,
                  left + Inches(0.65), top + Inches(0.55), col_w - Inches(1.0), Inches(0.55),
                  font_size=11, bold=False, color=C_LIGHT)


def _slide_executive_summary(prs, data, narratives):
    """Slide 3: Executive Summary"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Executive Summary")

    # KPI row
    kpis = [
        ("Client",    data.get("client_name", "-")),
        ("Industry",  data.get("industry", "-").replace("_", " ").title()),
        ("Timeline",  data.get("delivery_timeline", "-").replace("_", " ").title()),
        ("Budget",    f"${data.get('estimated_budget_usd', 0):,}"),
    ]
    for i, (label, val) in enumerate(kpis):
        _add_kpi_box(slide,
                     Inches(0.3) + i * Inches(3.2), Inches(1.2),
                     Inches(3.0), Inches(1.0),
                     label, val)

    _add_rect(slide, Inches(0.3), Inches(2.4), Inches(12.5), Inches(0.04), C_ACCENT)
    _add_text(slide, narratives.get("summary", ""),
              Inches(0.4), Inches(2.6), Inches(12.3), Inches(4.5),
              font_size=16, bold=False, color=C_DARK, align=PP_ALIGN.LEFT, wrap=True)


def _slide_problem(prs, data, narratives):
    """Slide 4: Problem Statement & Outcome"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Problem Statement & Outcome")

    # Left column – problem
    _add_rect(slide, Inches(0.3), Inches(1.3), Inches(6.0), Inches(5.8), C_MID)
    _add_text(slide, "🔴  Current Limitation",
              Inches(0.5), Inches(1.4), Inches(5.6), Inches(0.5),
              font_size=14, bold=True, color=C_ACCENT)
    _add_text(slide, narratives.get("problem", ""),
              Inches(0.5), Inches(1.9), Inches(5.6), Inches(5.0),
              font_size=14, bold=False, color=C_WHITE, wrap=True)

    # Right column – outcome
    _add_rect(slide, Inches(6.6), Inches(1.3), Inches(6.4), Inches(5.8), C_DARK)
    _add_text(slide, "🎯  Target Outcome",
              Inches(6.8), Inches(1.4), Inches(6.0), Inches(0.5),
              font_size=14, bold=True, color=C_ACCENT)
    facts = [
        f"Decision: {data.get('decision_supported', '-')}",
        f"Error Tolerance: {data.get('error_tolerance', '-').replace('_',' ').title()}",
        f"System Autonomy: {data.get('system_autonomy', '-').replace('_',' ').title()}",
        f"Optimization Goal: {data.get('optimization_goal', '-').replace('_',' ').title()}",
    ]
    for j, fact in enumerate(facts):
        _add_text(slide, f"▸  {fact}",
                  Inches(6.8), Inches(2.0) + j * Inches(0.7), Inches(6.0), Inches(0.6),
                  font_size=14, bold=False, color=C_LIGHT, wrap=True)


def _slide_data_reality(prs, data, narratives):
    """Slide 5: Data Reality"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Data Reality & Infrastructure")

    fields = [
        ("Data Source Type",  data.get("data_source_type",  "-").replace("_"," ").title()),
        ("Data Frequency",    data.get("data_frequency",    "-").replace("_"," ").title()),
        ("Data Quality",      data.get("data_quality",      "-").replace("_"," ").title()),
        ("Data Formats",      data.get("data_format",       ["-"]) if isinstance(data.get("data_format"), str) else ", ".join(data.get("data_format", ["-"]))),
        ("Data Volume",       f"{data.get('data_volume','-')} {data.get('data_volume_unit','')}".strip()),
        ("Inference Location",data.get("inference_location","-").replace("_"," ").title()),
        ("Cloud / Edge",      (data.get("cloud_provider") or data.get("edge_hardware") or "N/A").replace("_"," ").title()),
        ("Execution Trigger", data.get("execution_trigger","-").replace("_"," ").title()),
        ("Failure Handling",  data.get("failure_handling", "Human Review")),
    ]
    for i, (label, val) in enumerate(fields):
        col = i % 3
        row = i // 3
        _add_kpi_box(slide,
                     Inches(0.3) + col * Inches(4.3), Inches(1.3) + row * Inches(1.5),
                     Inches(4.1), Inches(1.3),
                     label, str(val))

    _add_text(slide, narratives.get("data", ""),
              Inches(0.3), Inches(6.1), Inches(12.5), Inches(1.2),
              font_size=12, bold=False, color=C_DARK, wrap=True)


def _slide_intelligence(prs, data, narratives):
    """Slide 6: Intelligence Layer"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Proposed Intelligence Layer")

    _add_text(slide, narratives.get("intelligence", ""),
              Inches(0.4), Inches(1.2), Inches(12.4), Inches(2.0),
              font_size=15, bold=False, color=C_DARK, wrap=True)

    cards = [
        ("System Autonomy",     data.get("system_autonomy", "-").replace("_"," ").title()),
        ("Optimization Goal",   data.get("optimization_goal", "-").replace("_"," ").title()),
        ("Retraining Cadence",  data.get("retraining_frequency", "-")),
        ("Execution Trigger",   data.get("execution_trigger", "-").replace("_"," ").title()),
    ]
    for i, (lbl, val) in enumerate(cards):
        _add_kpi_box(slide,
                     Inches(0.3) + i * Inches(3.2), Inches(3.5),
                     Inches(3.0), Inches(1.3),
                     lbl, val, bg=C_DARK)

    _add_text(slide, narratives.get("execution", ""),
              Inches(0.4), Inches(5.0), Inches(12.4), Inches(2.2),
              font_size=14, bold=False, color=C_DARK, wrap=True)


def _slide_architecture(prs, data):
    """Slide 7: Solution Architecture (diagram)"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_dark_bg(slide)
    _add_text(slide, "Solution Architecture",
              Inches(0.5), Inches(0.2), Inches(12.0), Inches(0.8),
              font_size=28, bold=True, color=C_ACCENT)
    _add_rect(slide, 0, Inches(1.0), SLIDE_W, Inches(0.04), C_ACCENT)

    arch = data.get("architecture_diagram_path")
    if arch:
        _add_image_safe(slide, arch, Inches(1.0), Inches(1.2),
                        width=Inches(11.0), height=Inches(5.8))
    else:
        _add_text(slide, "[Architecture Diagram — upload via the RFP form]",
                  Inches(2.0), Inches(3.5), Inches(9.0), Inches(1.0),
                  font_size=16, bold=False, color=C_LIGHT, align=PP_ALIGN.CENTER)


def _slide_security(prs, data):
    """Slide 8: Security & Compliance"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Security & Compliance")

    stds  = data.get("security_standards", [])
    ctrls = data.get("security_controls",  [])

    _add_rect(slide, Inches(0.3), Inches(1.3), Inches(6.0), Inches(5.8), C_MID)
    _add_text(slide, "Compliance Standards",
              Inches(0.5), Inches(1.4), Inches(5.5), Inches(0.5),
              font_size=14, bold=True, color=C_ACCENT)
    stds_text = "\n".join([f"✔  {s.replace('_',' ').upper()}" for s in stds]) if stds else "✔  Standard enterprise compliance"
    _add_text(slide, stds_text,
              Inches(0.5), Inches(2.0), Inches(5.5), Inches(5.0),
              font_size=14, bold=False, color=C_WHITE, wrap=True)

    _add_rect(slide, Inches(6.6), Inches(1.3), Inches(6.4), Inches(5.8), C_DARK)
    _add_text(slide, "Security Controls",
              Inches(6.8), Inches(1.4), Inches(6.0), Inches(0.5),
              font_size=14, bold=True, color=C_ACCENT)
    ctrls_text = "\n".join([f"🔒  {c.replace('_',' ').title()}" for c in ctrls]) if ctrls else "🔒  VPC, IAM, Encryption at rest"
    _add_text(slide, ctrls_text,
              Inches(6.8), Inches(2.0), Inches(6.0), Inches(5.0),
              font_size=14, bold=False, color=C_LIGHT, wrap=True)


def _slide_success_metrics(prs, data):
    """Slide 9: Success Metrics & Validation"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Success Metrics & Validation Approach")

    dims = data.get("success_dimensions", [])
    val_strat = data.get("validation_strategy", "Human-in-the-loop review.")

    # Dimension pills
    for i, dim in enumerate(dims[:6]):
        col = i % 3
        row = i // 3
        _add_kpi_box(slide,
                     Inches(0.3) + col * Inches(4.3), Inches(1.3) + row * Inches(1.5),
                     Inches(4.1), Inches(1.2),
                     "Success Dimension", dim.replace("_"," ").title(), bg=C_DARK)

    # Validation approach box
    top_y = Inches(4.5) if len(dims) > 3 else Inches(3.2)
    _add_rect(slide, Inches(0.3), top_y, Inches(12.5), Inches(0.04), C_ACCENT)
    _add_text(slide, "Validation Strategy",
              Inches(0.3), top_y + Inches(0.1), Inches(3.0), Inches(0.5),
              font_size=13, bold=True, color=C_MID)
    _add_text(slide, val_strat,
              Inches(0.3), top_y + Inches(0.6), Inches(12.5), Inches(2.5),
              font_size=14, bold=False, color=C_DARK, wrap=True)


def _slide_domain_deepdive(prs, data, proposal_type, narratives):
    """Slide 10: Domain-specific deep dive"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, f"{proposal_type} — Technical Deep Dive")

    if proposal_type == "Demand Forecasting":
        fields = [
            ("Forecasting Methods",  ", ".join([m.replace("_"," ").title() for m in data.get("forecasting_methods",[])])),
            ("Forecast Horizon",     f"{data.get('forecast_horizon_days','-')} days"),
            ("Historical Data",      f"{data.get('historical_data_years','-')} years"),
            ("Forecast Level",       data.get("forecast_level", "-")),
            ("Demand Volatility",    data.get("demand_volatility", "-")),
            ("Planning Decision",    data.get("planning_decision", "-")),
        ]
    elif proposal_type == "Visual Inspection":
        fields = [
            ("Inspection Type",  data.get("inspection_type","-").replace("_"," ").title()),
            ("Defect Categories", ", ".join(data.get("defect_categories",[]))),
            ("Accuracy Req.",    f"{data.get('accuracy_requirement',0)*100:.1f}%"),
            ("Image Source",     data.get("image_source","-").replace("_"," ").title()),
            ("Inspection Points",data.get("inspection_points","-")),
            ("Cost Matrix",      data.get("cost_matrix","-")),
        ]
    else:  # Chatbot
        fields = [
            ("Chatbot Type",        data.get("chatbot_type","-").replace("_"," ").title()),
            ("Queries / Day",       f"{data.get('expected_queries_per_day',0):,}"),
            ("Channel Strategy",    data.get("channel_strategy","-").replace("_"," ").title()),
            ("Orchestration",       data.get("orchestration_type","-").replace("_"," ").title()),
            ("Knowledge Strategy",  data.get("knowledge_strategy","-").replace("_"," ").title()),
            ("Guardrail Level",     data.get("guardrail_level","-").replace("_"," ").title()),
        ]

    for i, (lbl, val) in enumerate(fields):
        col = i % 3
        row = i // 3
        _add_kpi_box(slide,
                     Inches(0.3) + col * Inches(4.3), Inches(1.3) + row * Inches(1.5),
                     Inches(4.1), Inches(1.3),
                     lbl, str(val))

    _add_text(slide, narratives.get("technical", ""),
              Inches(0.4), Inches(6.0), Inches(12.4), Inches(1.3),
              font_size=12, bold=False, color=C_DARK, wrap=True)


def _slide_timeline(prs, data):
    """Slide 11: Implementation Timeline (Gantt table)"""
    timeline_data = data.get("timeline_data", [])
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_dark_bg(slide)
    _add_text(slide, "Implementation Timeline",
              Inches(0.5), Inches(0.2), Inches(12.0), Inches(0.8),
              font_size=28, bold=True, color=C_ACCENT)
    _add_rect(slide, 0, Inches(1.0), SLIDE_W, Inches(0.04), C_ACCENT)

    if not timeline_data:
        _add_text(slide, "No timeline data provided.",
                  Inches(0.5), Inches(3.5), Inches(12), Inches(1),
                  font_size=16, color=C_LIGHT)
        return

    cols = list(timeline_data[0].keys())
    rows = len(timeline_data)
    left   = Inches(0.3)
    top    = Inches(1.2)
    width  = Inches(12.7)
    height = Inches(5.8)

    tbl = slide.shapes.add_table(rows + 1, len(cols), left, top, width, height).table

    for i, col in enumerate(cols):
        cell = tbl.cell(0, i)
        cell.text = col
        cell.fill.solid()
        cell.fill.fore_color.rgb = C_MID
        for para in cell.text_frame.paragraphs:
            for run in para.runs:
                run.font.bold = True
                run.font.color.rgb = C_WHITE
                run.font.size = Pt(13)

    for r, row_data in enumerate(timeline_data):
        for c, col in enumerate(cols):
            val = str(row_data.get(col, ""))
            cell = tbl.cell(r + 1, c)
            cell.text = val
            if val.strip().upper() == "X":
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_GANTT
                cell.text = ""
            else:
                cell.fill.solid()
                cell.fill.fore_color.rgb = C_DARK if r % 2 == 0 else C_MID
            for para in cell.text_frame.paragraphs:
                for run in para.runs:
                    run.font.color.rgb = C_WHITE
                    run.font.size = Pt(12)


def _slide_commercials(prs, data):
    """Slide 12: Commercial Summary"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _set_light_bg(slide)
    _add_section_heading(slide, "Commercial Summary")

    budget   = f"${data.get('estimated_budget_usd', 0):,}"
    timeline = data.get("delivery_timeline", "-").replace("_"," ").title()
    vendor   = data.get("vendor_name", "Oneture")
    client   = data.get("client_name", "Client")

    kpis = [
        ("Total Budget Estimate", budget),
        ("Delivery Timeline",     timeline),
        ("Vendor",                vendor),
        ("Client",                client),
    ]
    for i, (lbl, val) in enumerate(kpis):
        _add_kpi_box(slide,
                     Inches(0.3) + i * Inches(3.2), Inches(1.3),
                     Inches(3.0), Inches(1.4),
                     lbl, val, bg=C_DARK)

    assumptions = data.get("key_assumptions", "").strip()
    if assumptions:
        _add_text(slide, "Key Assumptions",
                  Inches(0.4), Inches(3.0), Inches(6.0), Inches(0.5),
                  font_size=14, bold=True, color=C_MID)
        _add_text(slide, assumptions,
                  Inches(0.4), Inches(3.5), Inches(12.2), Inches(3.5),
                  font_size=14, bold=False, color=C_DARK, wrap=True)


def _slide_closing(prs, data):
    """Slide 13: Closing / Why Us"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    _add_rect(slide, 0, 0, SLIDE_W, SLIDE_H, C_DARK)
    _add_rect(slide, 0, Inches(4.3), SLIDE_W, Inches(0.06), C_ACCENT)

    vendor = data.get("vendor_name", "Oneture")
    _add_text(slide, "Why Choose Us",
              Inches(0.6), Inches(0.8), Inches(12.0), Inches(0.8),
              font_size=32, bold=True, color=C_ACCENT)
    _add_text(slide,
              f"{vendor} brings deep domain expertise in enterprise AI, proven delivery frameworks, "
              f"and a partnership-first approach. We combine cutting-edge ML with pragmatic engineering "
              f"to deliver measurable business outcomes — on time and within budget.",
              Inches(0.6), Inches(1.8), Inches(12.0), Inches(2.2),
              font_size=16, bold=False, color=C_LIGHT, wrap=True)

    bullets = ["Domain-Specific AI Expertise", "Agile, Milestone-Driven Delivery",
               "Secure & Scalable Infrastructure", "Transparent Pricing, No Hidden Costs"]
    for i, b in enumerate(bullets):
        _add_text(slide, f"✔  {b}",
                  Inches(0.6), Inches(4.5) + i * Inches(0.55), Inches(12.0), Inches(0.5),
                  font_size=16, bold=False, color=C_WHITE)

    _add_image_safe(slide, data.get("company_logo_path"),
                    Inches(11.3), Inches(6.2), height=Inches(1.0))


# ─── MAIN ENTRY POINT ──────────────────────────────────────────────────────────

def create_presentation(data: dict, proposal_type: str, api_key: str) -> io.BytesIO:
    """Builds a complete RFP proposal deck from scratch."""

    # 1. Generate AI narratives
    replacements = generate_narrative_content(data, proposal_type, api_key)

    # Parse the 6 narrative sections into a friendly dict
    narratives = {
        "summary":      replacements.get("{{EXECUTIVE_SUMMARY}}", ""),
        "technical":    replacements.get("{{TECH_NARRATIVE}}", ""),
        "problem":      replacements.get("{{PROBLEM_NARRATIVE}}", ""),
        "data":         replacements.get("{{DATA_NARRATIVE}}", ""),
        "intelligence": replacements.get("{{INTELLIGENCE_NARRATIVE}}", ""),
        "execution":    replacements.get("{{EXECUTION_NARRATIVE}}", ""),
    }

    # 2. Create blank presentation (widescreen 13.33 x 7.5)
    prs = Presentation()
    prs.slide_width  = SLIDE_W
    prs.slide_height = SLIDE_H

    # 3. Build slides in RFP order
    _slide_title(prs, data, proposal_type)
    _slide_agenda(prs, proposal_type)
    _slide_executive_summary(prs, data, narratives)
    _slide_problem(prs, data, narratives)
    _slide_data_reality(prs, data, narratives)
    _slide_intelligence(prs, data, narratives)
    _slide_architecture(prs, data)
    _slide_security(prs, data)
    _slide_success_metrics(prs, data)
    _slide_domain_deepdive(prs, data, proposal_type, narratives)
    _slide_timeline(prs, data)
    _slide_commercials(prs, data)
    _slide_closing(prs, data)

    # 4. Save to memory
    output = io.BytesIO()
    prs.save(output)
    output.seek(0)
    return output
