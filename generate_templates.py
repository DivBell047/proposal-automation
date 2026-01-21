# Cell 2: The Template Factory (Updated)
# Run this ONCE to regenerate templates with Client/Vendor placeholders.

import os
from pptx import Presentation
from pptx.util import Inches, Pt

# 1. Create directory
if not os.path.exists("templates"):
    os.makedirs("templates")

def create_base_template(filename, title_pattern, specific_placeholders):
    prs = Presentation()

    # --- SLIDE 1: Title Slide (UPDATED) ---
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = title_pattern
    # Added {{CLIENT}} and {{VENDOR}} placeholders here
    slide.placeholders[1].text = "Prepared for: {{CLIENT}}\nPresented by: {{VENDOR}}\nIndustry: {{INDUSTRY}}"

    # --- SLIDE 2: Executive Summary ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Executive Summary"
    slide.placeholders[1].text = "{{EXECUTIVE_SUMMARY}}"

    # --- SLIDE 3: Solution Scope ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Proposed Solution Scope"
    tf = slide.placeholders[1].text_frame
    tf.text = "Technical Approach:"
    for ph in specific_placeholders:
        p = tf.add_paragraph()
        p.text = ph
        p.level = 0

    # --- SLIDE 4: Commercials ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Commercials & Timeline"
    tf = slide.placeholders[1].text_frame
    tf.text = "Project Details"
    
    # Simple lines for budget/timeline
    p = tf.add_paragraph()
    p.text = "Budget: {{BUDGET}}"
    p = tf.add_paragraph()
    p.text = "Timeline: {{TIMELINE}}"
    p = tf.add_paragraph()
    p = tf.add_paragraph()
    p.text = "Cloud: {{CLOUD}}"

    # --- SLIDE 5: Problem & Outcome ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Problem Statement & Outcome"
    slide.placeholders[1].text = "{{PROBLEM_NARRATIVE}}"

    # --- SLIDE 6: Data Reality ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Data Reality"
    slide.placeholders[1].text = "{{DATA_NARRATIVE}}"

    # --- SLIDE 7: Intelligence Layer ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Intelligence Layer"
    slide.placeholders[1].text = "{{INTELLIGENCE_NARRATIVE}}"

    # --- SLIDE 8: Execution & Deployment ---
    slide = prs.slides.add_slide(prs.slide_layouts[1])
    slide.shapes.title.text = "Execution Strategy"
    slide.placeholders[1].text = "{{EXECUTION_NARRATIVE}}"

    save_path = os.path.join("templates", filename)
    prs.save(save_path)
    print(f"✅ Created updated template: {save_path}")

# Generate the 3 required templates
create_base_template("demand_forecasting.pptx", "Demand Forecasting", 
    ["Method: {{METHODS}}", "Horizon: {{HORIZON}}", "Drivers: {{FEATURES_NARRATIVE}}"])

create_base_template("visual_inspection.pptx", "Visual Inspection System", 
    ["Type: {{INSPECTION_TYPE}}", "Defects: {{DEFECTS}}", "Accuracy: {{ACCURACY}}"])

create_base_template("chatbot.pptx", "AI Chatbot Solution", 
    ["Type: {{CHATBOT_TYPE}}", "Platforms: {{PLATFORMS}}", "Volume: {{QUERIES}}"])